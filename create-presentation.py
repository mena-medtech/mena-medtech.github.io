from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

# Create presentation
prs = Presentation()
prs.slide_width = Inches(10)
prs.slide_height = Inches(7.5)

# Define colors
PRIMARY_COLOR = RGBColor(102, 126, 234)  # Purple
ACCENT_COLOR = RGBColor(245, 87, 108)   # Pink
WHITE = RGBColor(255, 255, 255)
DARK_GRAY = RGBColor(51, 51, 51)
LIGHT_GRAY = RGBColor(248, 249, 250)

def add_title_slide(prs, title, subtitle):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = PRIMARY_COLOR
    
    # Add title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(2.5), Inches(9), Inches(1.5))
    title_frame = title_box.text_frame
    title_frame.word_wrap = True
    p = title_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(60)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER
    
    # Add subtitle
    subtitle_box = slide.shapes.add_textbox(Inches(0.5), Inches(4.2), Inches(9), Inches(1.5))
    subtitle_frame = subtitle_box.text_frame
    subtitle_frame.word_wrap = True
    p = subtitle_frame.paragraphs[0]
    p.text = subtitle
    p.font.size = Pt(28)
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER
    
    return slide

def add_content_slide(prs, title, bg_color=WHITE):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = bg_color
    
    # Add title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(9), Inches(0.8))
    title_frame = title_box.text_frame
    p = title_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_COLOR if bg_color == WHITE else WHITE
    
    return slide

# Slide 1: Title Slide
add_title_slide(prs, "MENA MedTech", "Medical Device Engineering & Regulatory Expertise\n15 Years of Excellence")

# Slide 2: About
slide = add_content_slide(prs, "About MENA MedTech")
content = [
    ("15+", "Years of Experience"),
    ("100+", "Projects Completed"),
    ("50+", "Devices Commercialized")
]
y_pos = 1.5
for stat, label in content:
    stat_box = slide.shapes.add_textbox(Inches(1 + (content.index((stat, label)) * 3)), Inches(y_pos), Inches(2.5), Inches(0.6))
    stat_frame = stat_box.text_frame
    p = stat_frame.paragraphs[0]
    p.text = stat
    p.font.size = Pt(48)
    p.font.bold = True
    p.font.color.rgb = ACCENT_COLOR
    p.alignment = PP_ALIGN.CENTER
    
    label_box = slide.shapes.add_textbox(Inches(1 + (content.index((stat, label)) * 3)), Inches(y_pos + 0.7), Inches(2.5), Inches(0.4))
    label_frame = label_box.text_frame
    p = label_frame.paragraphs[0]
    p.text = label
    p.font.size = Pt(16)
    p.font.color.rgb = DARK_GRAY
    p.alignment = PP_ALIGN.CENTER

description_box = slide.shapes.add_textbox(Inches(0.5), Inches(3.2), Inches(9), Inches(3.5))
description_frame = description_box.text_frame
description_frame.word_wrap = True
p = description_frame.paragraphs[0]
p.text = "With 15 years of proven expertise in medical device engineering and regulatory compliance, MENA MedTech delivers comprehensive solutions for the development and commercialization of innovative medical devices."
p.font.size = Pt(18)
p.font.color.rgb = DARK_GRAY
p.alignment = PP_ALIGN.CENTER

# Slide 3: Expertise Areas
slide = add_content_slide(prs, "Our Expertise Areas")
expertise = [
    ("Electromechanical Devices", ["Motor & actuator integration", "Control systems design", "Prototype development"]),
    ("Combination Products", ["Drug-device strategy", "Dual regulatory pathways", "Coordinated commercialization"]),
    ("Regulatory Strategy", ["FDA pathway selection", "CE Mark compliance", "Market entry strategy"])
]

y_pos = 1.5
for i, (title, items) in enumerate(expertise):
    x_pos = 0.5 + (i * 3.2)
    title_box = slide.shapes.add_textbox(Inches(x_pos), Inches(y_pos), Inches(2.8), Inches(0.5))
    title_frame = title_box.text_frame
    p = title_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_COLOR
    
    content_box = slide.shapes.add_textbox(Inches(x_pos), Inches(y_pos + 0.6), Inches(2.8), Inches(5))
    content_frame = content_box.text_frame
    content_frame.word_wrap = True
    for item in items:
        p = content_frame.add_paragraph()
        p.text = "• " + item
        p.font.size = Pt(11)
        p.font.color.rgb = DARK_GRAY
        p.level = 0

# Slide 4: Core Services
slide = add_content_slide(prs, "Core Services", PRIMARY_COLOR)
services = [
    "Project Charter", "Design & Development Plan", "Design Control Phases",
    "R&D Engineering", "Manufacturing Engineering", "Quality Engineering",
    "Design Transfer", "Post-Market Surveillance", "Technical File & Registration"
]

y_start = 1.5
x_start = 0.5
cols = 3
rows = 3

for idx, service in enumerate(services):
    col = idx % cols
    row = idx // cols
    x = x_start + (col * 3.1)
    y = y_start + (row * 1.8)
    
    service_box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(2.8), Inches(1.5))
    service_frame = service_box.text_frame
    service_frame.word_wrap = True
    p = service_frame.paragraphs[0]
    p.text = service
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER

# Slide 5: Consulting Services
slide = add_content_slide(prs, "Consulting Services")
consulting = [
    ("Project Management", "Planning, scheduling, resource allocation"),
    ("Strategic Planning", "Concept evaluation, market assessment, roadmap"),
    ("Regulatory & Compliance", "FDA/CE strategy, compliance planning, submissions"),
    ("Technical & Manufacturing", "Design optimization, process planning, scale-up"),
    ("Combination Products", "Drug-device strategy, interaction studies"),
    ("Market Entry & Expansion", "Market analysis, business strategy, launch planning")
]

y_pos = 1.5
for i, (title, desc) in enumerate(consulting):
    x_pos = 0.5 if i % 2 == 0 else 5.2
    if i % 2 == 0 and i > 0:
        y_pos += 1.8
    
    title_box = slide.shapes.add_textbox(Inches(x_pos), Inches(y_pos), Inches(4.3), Inches(0.4))
    title_frame = title_box.text_frame
    p = title_frame.paragraphs[0]
    p.text = title
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_COLOR
    
    desc_box = slide.shapes.add_textbox(Inches(x_pos), Inches(y_pos + 0.45), Inches(4.3), Inches(1.2))
    desc_frame = desc_box.text_frame
    desc_frame.word_wrap = True
    p = desc_frame.paragraphs[0]
    p.text = desc
    p.font.size = Pt(11)
    p.font.color.rgb = DARK_GRAY

# Slide 6: Pricing - Gulf Region
slide = add_content_slide(prs, "Service Pricing - Gulf Region (SAR)", ACCENT_COLOR)
pricing_gulf = [
    ("Project Charter", "45,000"),
    ("Design & Dev Plan", "75,000"),
    ("Design Control", "150,000"),
    ("R&D Engineering", "250,000"),
    ("Manufacturing Eng", "200,000"),
    ("Quality Engineering", "120,000")
]

y_pos = 1.5
for i, (service, price) in enumerate(pricing_gulf):
    col = i % 3
    row = i // 3
    x = 0.5 + (col * 3.1)
    y = y_pos + (row * 2.5)
    
    service_box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(2.8), Inches(0.4))
    service_frame = service_box.text_frame
    p = service_frame.paragraphs[0]
    p.text = service
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER
    
    price_box = slide.shapes.add_textbox(Inches(x), Inches(y + 0.5), Inches(2.8), Inches(0.5))
    price_frame = price_box.text_frame
    p = price_frame.paragraphs[0]
    p.text = price
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.alignment = PP_ALIGN.CENTER

# Slide 7: Consulting Pricing
slide = add_content_slide(prs, "Consulting Services Pricing (Monthly Rates)")
consulting_pricing = [
    ("Project Management", "SAR 25K-40K", "USD 7K-11K"),
    ("Strategic Planning", "SAR 30K-50K", "USD 8.5K-14K"),
    ("Regulatory & Compliance", "SAR 35K-55K", "USD 10K-16K"),
    ("Combination Products", "SAR 40K-60K", "USD 12K-18K")
]

y_pos = 1.6
for i, (service, gulf, us) in enumerate(consulting_pricing):
    x_pos = 0.5 if i % 2 == 0 else 5.3
    if i % 2 == 0 and i > 0:
        y_pos += 2
    
    service_box = slide.shapes.add_textbox(Inches(x_pos), Inches(y_pos), Inches(4.2), Inches(0.4))
    service_frame = service_box.text_frame
    p = service_frame.paragraphs[0]
    p.text = service
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_COLOR
    
    pricing_box = slide.shapes.add_textbox(Inches(x_pos), Inches(y_pos + 0.45), Inches(4.2), Inches(1.2))
    pricing_frame = pricing_box.text_frame
    pricing_frame.word_wrap = True
    p = pricing_frame.paragraphs[0]
    p.text = f"Gulf Region: {gulf}\nUnited States: {us}"
    p.font.size = Pt(11)
    p.font.color.rgb = DARK_GRAY

# Slide 8: Why Choose Us
slide = add_content_slide(prs, "Why Choose MENA MedTech")
features = [
    "15+ years proven track record",
    "50+ commercialized devices",
    "Multi-regional expertise",
    "FDA & CE Mark proficiency",
    "End-to-end solutions",
    "Combination product specialists",
    "Regulatory excellence",
    "Quality-focused approach"
]

y_pos = 1.6
for i, feature in enumerate(features):
    x_pos = 0.7 if i < 4 else 5.3
    y = y_pos + ((i % 4) * 1.2)
    
    feature_box = slide.shapes.add_textbox(Inches(x_pos), Inches(y), Inches(4), Inches(0.8))
    feature_frame = feature_box.text_frame
    feature_frame.word_wrap = True
    p = feature_frame.paragraphs[0]
    p.text = "✓ " + feature
    p.font.size = Pt(14)
    p.font.color.rgb = PRIMARY_COLOR
    p.font.bold = True

# Slide 9: Our Process
slide = add_content_slide(prs, "Our Comprehensive Process", PRIMARY_COLOR)
process_steps = [
    ("Discovery & Planning", ["Concept evaluation", "Market assessment", "Technical feasibility"]),
    ("Development & Design", ["Design control", "Prototype development", "Testing & validation"]),
    ("Manufacturing & Scale", ["Process optimization", "Design transfer", "Production scale-up"]),
    ("Regulatory & Launch", ["Regulatory submission", "Compliance verification", "Market launch"])
]

y_pos = 1.8
for i, (step, items) in enumerate(process_steps):
    x_pos = 0.5 + (i % 2) * 4.9
    if i == 2:
        y_pos = 4.2
    
    step_box = slide.shapes.add_textbox(Inches(x_pos), Inches(y_pos), Inches(4.3), Inches(0.5))
    step_frame = step_box.text_frame
    p = step_frame.paragraphs[0]
    p.text = f"Step {i+1}: {step}"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = WHITE
    
    items_box = slide.shapes.add_textbox(Inches(x_pos), Inches(y_pos + 0.55), Inches(4.3), Inches(1.5))
    items_frame = items_box.text_frame
    items_frame.word_wrap = True
    for item in items:
        p = items_frame.add_paragraph()
        p.text = "• " + item
        p.font.size = Pt(10)
        p.font.color.rgb = WHITE

# Slide 10: Contact
slide = add_content_slide(prs, "Get in Touch", ACCENT_COLOR)

contact_info = [
    ("WhatsApp", "+1 508 410 4492"),
    ("Email", "services@mena-medtech.com"),
    ("Website", "www.mena-medtech.com")
]

y_pos = 2
for i, (label, value) in enumerate(contact_info):
    label_box = slide.shapes.add_textbox(Inches(2), Inches(y_pos + (i * 1.2)), Inches(6), Inches(0.4))
    label_frame = label_box.text_frame
    p = label_frame.paragraphs[0]
    p.text = label
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = WHITE
    
    value_box = slide.shapes.add_textbox(Inches(2), Inches(y_pos + 0.4 + (i * 1.2)), Inches(6), Inches(0.4))
    value_frame = value_box.text_frame
    p = value_frame.paragraphs[0]
    p.text = value
    p.font.size = Pt(14)
    p.font.color.rgb = WHITE

# Slide 11: Thank You
add_title_slide(prs, "Thank You", "Your Partner in Medical Device Excellence\n© 2024 MENA MedTech")

# Save presentation
prs.save('MENA-MedTech-Services-Presentation.pptx')
print("Presentation created successfully: MENA-MedTech-Services-Presentation.pptx")